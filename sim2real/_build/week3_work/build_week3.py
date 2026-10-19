from pathlib import Path
from hashlib import sha256
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

work = Path(__file__).resolve().parent
sim2real = work.parents[1]
template = sim2real.parents[1] / 'weekly_report' / 'Week2_Report.docx'
sections = [
    ('Week 3 progress', [
        'This week I reviewed the sim-to-real literature and connected its findings to DroneRanger\'s hardware plan. Building on last week\'s flight-stack review, I focused on what simulation can tell us about a real garden patrol. The main lesson is that sharing a PX4 and ROS 2 interface makes testing reusable, but reliable transfer still requires calibration and measurements. The literature review provides a practical basis for deciding what to test before outdoor flight.'
    ]),
    ('Findings relevant to DroneRanger', [
        'The Spring prototype exposed several gaps: roughly 0.5 m of altitude error near the ground, waypoints reported as reached despite visible position offsets, and invalid depth images. These issues matter directly to the hardware group\'s goal of a reliable Mission-mode garden patrol. Ideal sensor readings and simplified dynamics in simulation can hide the problems that appear during takeoff, hovering and navigation outdoors.',
        'Neural Lander shows why near-ground aerodynamics deserve attention. Ground effect is a possible contributor to altitude error, although the prototype\'s cause still needs diagnosis. Work on quadrotor system identification shows how flight logs can help estimate thrust and motor delay. Swift similarly illustrates how measured errors can improve a simulator. For DroneRanger, these studies support measuring the aircraft and updating the X500 model before varying uncertain parameters across simulated tests.',
        'The perception literature also explains why perfect simulated depth is insufficient. The OAK-D Lite can produce invalid pixels and misleading distance readings. Learning High-Speed Flight in the Wild demonstrates the value of training with realistic stereo-depth artifacts. Our hardware contribution should therefore include depth recordings under different lighting and distances, with invalid readings treated as unknown. Neural-Fly highlights wind as another outdoor gap, supporting rehearsals under varied wind conditions and logging conditions during real tests.'
    ]),
    ('Implications for RL cooperation', [
        'The RL environment uses a simplified controller and observations that do not yet match the real platform. A shared interface should specify available sensors, units, coordinate frames, action meanings, timing and invalid-depth handling. The policy should first be evaluated through PX4 SITL, with PX4 retaining flight stabilization. Hardware measurements of sensing and communication delays will help the RL group test realistic inputs. Policy-driven Offboard flight remains a later extension beyond the current Mission-mode patrol milestone.'
    ]),
    ('Week 4 plan', [
        'Next week I plan to turn the review into a short test checklist and continue diagnosing the PX4-to-Jetson connection with propellers removed. I will define paired simulation and flight measurements for hover error, waypoint error and failsafe response, using matched firmware, mission files and parameters. Altitude and position should be checked against independent references rather than only the onboard estimate.',
        'After the hardware checks and flight approvals, repeated trials can establish a baseline and guide simulator updates. Each run should retain flight logs, parameters and environmental conditions. The planned altitude target is an RMSE of 0.15 m or less. A garden digital twin is useful future work; the immediate deliverable is a testable, calibrated platform for the patrol and later RL integration.'
    ])
]

doc = Document(template)
for element in list(doc.element.body):
    if element.tag != qn('w:sectPr'):
        doc.element.body.remove(element)
for rel_id, rel in list(doc.part.rels.items()):
    if rel.reltype == RT.HYPERLINK:
        del doc.part.rels[rel_id]
for name in ('Title', 'Heading 1', 'Normal'):
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.add_paragraph('DroneRanger Week 3 Report', 'Title')
p = doc.add_paragraph('Jiarui Zha, zha.j@northeastern.edu')
p.paragraph_format.space_after = Pt(7)
for heading, paragraphs in sections:
    doc.add_paragraph(heading, 'Heading 1')
    for text in paragraphs:
        doc.add_paragraph(text)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(0)
r = p.add_run('Sources: ')
r.font.size = Pt(8.5)
for i, (label, source) in enumerate([
    ('Literature Review V2', sim2real / 'DroneRanger_Literature_Review_V2.docx'),
    ('Comprehensive Sim2Real Review', sim2real / 'Sim2Real_Literature_Review_Comprehensive.md'),
]):
    if i:
        p.add_run(' | ').font.size = Pt(8.5)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), p.part.relate_to(source.as_uri(), RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '17'); props.append(size)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '333333'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    hyperlink.append(run); p._p.append(hyperlink)
doc.core_properties.title = 'DroneRanger Week 3 Report'
doc.core_properties.subject = 'Literature findings and sim-to-real validation plan'
doc.core_properties.author = 'Jiarui Zha'
out = work / 'Week3_Report.docx'
doc.save(out)
body = '\n\n'.join(text for _, paragraphs in sections for text in paragraphs)
(work / 'report_text.txt').write_text(body, encoding='utf-8')
print(f'Body words: {len(body.split())}')
print(f'Output: {out}')
print(f'Template SHA256: {sha256(template.read_bytes()).hexdigest()}')
