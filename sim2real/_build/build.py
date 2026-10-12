"""Build the sim2real deliverables from one reference database.

    python sim2real/_build/build.py

Inputs (this folder)
  refs.py                  reference database (verified metadata, LibKey ids)
  part*.src.md             comprehensive review in parts (concatenated in file-name order);
                           citations written as [@key] or [@key1; @key2]
  brief.py                 content of the brief (Canvas) version

Outputs (sim2real/)
  Sim2Real_Literature_Review_Comprehensive.md
  references.md            numbered IEEE list (same numbers as the comprehensive review) with access links
  references.bib           BibTeX for every cited reference (import into Zotero: File > Import)
  DroneRanger_Sim2Real_Literature_Review_Brief.docx
  references_brief.bib     BibTeX for the brief's references only
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
sys.dont_write_bytecode = True   # keep _build/ free of __pycache__
sys.path.insert(0, str(HERE))
from refs import LIBKEY_URL, REFS  # noqa: E402

BY_KEY = {r['key']: r for r in REFS}
assert len(BY_KEY) == len(REFS), 'duplicate reference keys'

MONTHS = ['Jan.', 'Feb.', 'Mar.', 'Apr.', 'May', 'Jun.', 'Jul.', 'Aug.', 'Sep.', 'Oct.', 'Nov.', 'Dec.']
BIB_MONTHS = ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec']
KIND_LABEL = {'J': 'Journal article', 'C': 'Conference paper', 'CH': 'Book chapter', 'B': 'Book',
              'P': 'Preprint (not peer reviewed)', 'D': 'Documentation', 'R': 'Report (internal)',
              'W': 'Software repository'}
KIND_PLURAL = {'J': 'journal articles', 'C': 'conference papers', 'CH': 'book chapters', 'B': 'books',
               'P': 'preprints', 'D': 'documentation pages', 'R': 'internal reports', 'W': 'software repositories'}
CITE = re.compile(r'\[@([^\]]+)\]')


# ------------------------------------------------------------------------------------------- IEEE formatting
def _initials(given):
    out = []
    for part in given.split():
        out.append('-'.join(s[0] + '.' for s in part.split('-') if s))
    return ' '.join(out)


def _person(name):
    if ',' not in name:
        return name                                   # organisation
    family, given = (s.strip() for s in name.split(',', 1))
    return f'{_initials(given)} {family}'.strip()


def ieee_authors(r):
    names = r['authors']
    persons = [n for n in names if ',' in n]
    if persons and ',' not in names[0]:              # 'OpenAI', 'NVIDIA' lead a list of people
        names = persons
    if len(names) > 6 or r.get('et_al'):
        return f'{_person(names[0])} <i>et al.</i>'
    names = [_person(n) for n in names]
    if len(names) == 1:
        return names[0]
    if len(names) == 2:
        return f'{names[0]} and {names[1]}'
    return ', '.join(names[:-1]) + ', and ' + names[-1]


def _date(r):
    return f"{MONTHS[r['month'] - 1]} {r['year']}" if r.get('month') else str(r['year'])


def _pages(r):
    if r.get('article_no'):
        return f"Art. no. {r['article_no']}"
    if r.get('pages'):
        p = r['pages'].replace('--', '–')
        return f'pp. {p}' if '–' in p else f'p. {p}'
    return None


def _accessed(iso):
    y, m, d = (int(x) for x in iso.split('-'))
    return f'{MONTHS[m - 1]} {d}, {y}'


def _q(title):
    """Quoted title with IEEE punctuation: no comma after a closing ? or !."""
    return f'\u201c{title}\u201d' if title[-1] in '?!' else f'\u201c{title},\u201d'


def ieee(r):
    """IEEE reference string with <i> markup (no trailing number)."""
    k, a, t = r['kind'], ieee_authors(r), r['title']
    doi = f", doi: {r['doi']}" if r.get('doi') else ''
    if k == 'J':
        parts = [f"<i>{r['venue']}</i>"]
        if r.get('volume'):
            parts.append(f"vol. {r['volume']}")
        if r.get('number'):
            parts.append(f"no. {r['number'].replace('--', '–')}")
        if _pages(r):
            parts.append(_pages(r))
        parts.append(_date(r))
        return f'{a}, {_q(t)} ' + ', '.join(parts) + doi + '.'
    if k == 'C':
        s = f"{a}, {_q(t)} in <i>{r['venue']}</i>"
        if r.get('editors'):
            eds = [_person(e) for e in r['editors']]
            s += ', ' + ' and '.join(eds) + (', Eds.' if len(eds) > 1 else ', Ed.')
            s += f" {r['address']}: {r['publisher']}, {r['year']}"
        elif r.get('series'):
            s += f", ser. {r['series']}, vol. {r['volume']}, {r['year']}"
        else:
            s += f", {r['year']}"
        if _pages(r):
            s += f', {_pages(r)}'
        if r.get('doi'):
            return s + doi + '.'
        if r.get('url'):
            return s + f". [Online]. Available: {r['url']}"
        if r.get('arxiv'):
            return s + f", arXiv:{r['arxiv']}."
        return s + '.'
    if k == 'CH':
        s = f"{a}, {_q(t)} in <i>{r['venue']}</i>"
        if r.get('editors'):
            eds = [_person(e) for e in r['editors']]
            s += ', ' + ' and '.join(eds) + (', Eds.' if len(eds) > 1 else ', Ed.')
        if r.get('series'):
            s += f", ser. {r['series']}, vol. {r['volume']}"
        s += f". {r['address']}: {r['publisher']}, {r['year']}"
        if _pages(r):
            s += f', {_pages(r)}'
        return s + doi + '.'
    if k == 'B':
        s = f"{a}, <i>{t}</i>"
        s += f", {r['edition']} ed." if r.get('edition') else '.'
        return s + f" {r['address']}: {r['publisher']}, {r['year']}" + doi + '.'
    if k == 'P':
        s = f"{a}, {_q(t)} {_date(r)}, arXiv:{r['arxiv']}"
        return s + (f". {r['note']}." if r.get('note') else '.')
    if k == 'D':
        return (f"{a}, {_q(t)} <i>{r['site']}</i>. Accessed: {_accessed(r['accessed'])}. "
                f"[Online]. Available: {r['url']}")
    if k == 'R':
        return f"{a}, {_q(t)} {r['institution']}, {r['address']}, {r['rtype']}, {_date(r)}."
    if k == 'W':
        return f"{a}, {_q(t)} {r['note']}, {_date(r)}. [Online]. Available: {r['url']}"
    raise ValueError(f"unknown kind {k} for {r['key']}")


def access_links_md(r):
    """Markdown access links: LibKey (journal articles), DOI, arXiv or URL."""
    links = []
    if r.get('libkey'):
        lid, ltype, oa = r['libkey']
        path = 'full-text-file' if ltype == 'fullTextFile' else 'content-location'
        label = '\U0001F4C4 **Download PDF**' if ltype == 'fullTextFile' else '\U0001F517 **View Article**'
        links.append(f"[{label}]({LIBKEY_URL.format(id=lid, path=path)})" + (' (Open Access)' if oa else ''))
    if r.get('doi'):
        links.append(f"[doi:{r['doi']}](https://doi.org/{r['doi']})")
    if r.get('arxiv'):
        links.append(f"[arXiv:{r['arxiv']}](https://arxiv.org/abs/{r['arxiv']})")
    if r.get('url') and not r.get('doi'):
        links.append(f"[link]({r['url']})")
    return ' · '.join(links)


def md_inline(s):
    return s.replace('<i>', '*').replace('</i>', '*').replace('<b>', '**').replace('</b>', '**')


# ------------------------------------------------------------------------------------------- citations
class Numbering:
    def __init__(self):
        self.order = []

    def number(self, key):
        if key not in BY_KEY:
            raise KeyError(f'unknown citation key: {key}')
        if key not in self.order:
            self.order.append(key)
        return self.order.index(key) + 1

    def sub(self, text, link=False):
        def repl(m):
            keys = [k.strip().lstrip('@') for k in m.group(1).split(';')]
            nums = sorted({self.number(k) for k in keys})
            parts, i = [], 0
            while i < len(nums):                       # IEEE: [3]-[5] for runs of three or more
                j = i
                while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                    j += 1
                if j - i >= 2:
                    parts.append(f'[{nums[i]}]–[{nums[j]}]')
                else:
                    parts.extend(f'[{n}]' for n in nums[i:j + 1])
                i = j + 1
            out = ', '.join(parts)
            if link:                                   # make [n] clickable in Markdown
                out = re.sub(r'\[(\d+)\]', r'[[\1]](#ref-\1)', out)
            return out
        return CITE.sub(repl, text)


# ------------------------------------------------------------------------------------------- BibTeX
def _bib_name(n):
    return n if ',' in n else '{' + n + '}'


def bibtex(r):
    k = r['kind']
    etype = {'J': 'article', 'C': 'inproceedings', 'CH': 'incollection', 'B': 'book', 'P': 'misc',
             'D': 'misc', 'R': 'techreport', 'W': 'misc'}[k]
    f = {}
    f['author'] = ' and '.join(_bib_name(n) for n in r['authors']) + (' and others' if r.get('et_al') else '')
    f['title'] = '{' + r['title'] + '}'
    if k == 'J':
        f['journal'] = r['venue_full']
        for src, dst in (('volume', 'volume'), ('number', 'number')):
            if r.get(src):
                f[dst] = r[src]
        if r.get('pages') or r.get('article_no'):
            f['pages'] = r.get('pages') or r['article_no']
    elif k in ('C', 'CH'):
        f['booktitle'] = r.get('venue_full', r['venue'])
        if r.get('editors'):
            f['editor'] = ' and '.join(r['editors'])
        for fld in ('series', 'volume', 'number', 'pages', 'publisher', 'address'):
            if r.get(fld):
                f[fld] = r[fld]
    elif k == 'B':
        for fld in ('edition', 'publisher', 'address', 'isbn'):
            if r.get(fld):
                f[fld] = r[fld]
    elif k == 'P':
        f['eprint'] = r['arxiv']
        f['archiveprefix'] = 'arXiv'
        f['url'] = f"https://arxiv.org/abs/{r['arxiv']}"
    elif k == 'D':
        f['howpublished'] = r['site']
        f['url'] = r['url']
        f['urldate'] = r['accessed']
        f['note'] = f"Accessed: {_accessed(r['accessed'])}"
    elif k == 'R':
        f['institution'] = r['institution']
        f['address'] = r['address']
        f['type'] = r['rtype']
    elif k == 'W':
        f['howpublished'] = r['note']
        f['url'] = r['url']
    if r.get('year') and k != 'D':
        f['year'] = str(r['year'])
    if r.get('month') and k != 'D':
        f['month'] = BIB_MONTHS[r['month'] - 1]
    if r.get('doi'):
        f['doi'] = r['doi']
    if r.get('url') and 'url' not in f and not r.get('doi'):
        f['url'] = r['url']
    if r.get('note') and k not in ('W',) and 'note' not in f:
        f['note'] = r['note']
    if r.get('arxiv') and k == 'C':
        f['eprint'] = r['arxiv']
        f['archiveprefix'] = 'arXiv'
    lines = [f'@{etype}{{{r["key"]},']
    for fld, val in f.items():
        val = val if fld == 'month' else '{' + val + '}'
        lines.append(f'  {fld:<14}= {val},')
    lines.append('}')
    return '\n'.join(lines)


def write_bib(path, keys, header):
    body = '\n\n'.join(bibtex(BY_KEY[k]) for k in keys)
    path.write_text(header + '\n\n' + body + '\n', encoding='utf-8')


# ------------------------------------------------------------------------------------------- comprehensive (Markdown)
def review_source():
    parts = sorted(HERE.glob('part*.src.md'))
    if not parts:
        raise SystemExit('no part*.src.md files found')
    src = '\n\n'.join(part.read_text(encoding='utf-8').strip() for part in parts) + '\n'
    left = re.findall(r'\{\{[A-Z0-9_]+\}\}', src)
    if left:
        raise SystemExit(f'unfilled placeholders in the review: {sorted(set(left))}')
    return src


def build_comprehensive():
    src = review_source()
    num = Numbering()
    body = num.sub(src, link=True)
    lines = ['', '## References', '',
             '*IEEE style, numbered in order of first citation. Journal articles were resolved through LibKey '
             '(Northeastern University Library); conference papers, chapters and books were checked against '
             'Crossref, PMLR or arXiv. Access links, peer-review status and BibTeX are in '
             '[references.md](references.md) and [references.bib](references.bib).*', '']
    for n, key in enumerate(num.order, 1):
        lines.append(f'<a id="ref-{n}"></a>[{n}] {md_inline(ieee(BY_KEY[key]))}  ')
    (OUT / 'Sim2Real_Literature_Review_Comprehensive.md').write_text(body.rstrip() + '\n' + '\n'.join(lines) + '\n',
                                                                     encoding='utf-8')
    return num.order


def build_reference_md(order, brief_keys):
    counts = {}
    for k in order:
        counts[BY_KEY[k]['kind']] = counts.get(BY_KEY[k]['kind'], 0) + 1
    summary = ', '.join(f'{v} {KIND_PLURAL[k]}' for k, v in sorted(counts.items(), key=lambda x: -x[1]))
    out = ['# Sim2Real reference list', '',
           f'{len(order)} references, numbered as in '
           '[Sim2Real_Literature_Review_Comprehensive.md](Sim2Real_Literature_Review_Comprehensive.md) '
           '(IEEE style, order of first citation). '
           f'Breakdown: {summary}.', '',
           '- **Verification (2026-10-11).** Every journal article was resolved through LibKey (Northeastern '
           'University Library); the \U0001F4C4 links open the library full text (sign-in may be required). '
           'LibKey does not index conference papers, chapters or books, so their DOIs were checked on Crossref '
           '(PMLR papers on proceedings.mlr.press, preprints on arXiv).',
           '- **Peer review.** Entries marked *Preprint* had not completed peer review when checked; '
           '*Documentation* entries describe system behaviour, not research results.',
           '- **Brief version.** ★ marks the references used in '
           '[DroneRanger_Sim2Real_Literature_Review_Brief.docx](DroneRanger_Sim2Real_Literature_Review_Brief.docx) '
           '(BibTeX subset: [references_brief.bib](references_brief.bib)).',
           '- **Citation manager.** [references.bib](references.bib) holds all entries; import it into Zotero '
           'with *File → Import…*, or use it directly with BibTeX/biblatex.', '']
    for n, key in enumerate(order, 1):
        r = BY_KEY[key]
        star = ' ★' if key in brief_keys else ''
        out.append(f'**[{n}]**{star} {md_inline(ieee(r))}  ')
        links = access_links_md(r) or {'R': 'internal project document (not public)',
                                        'C': 'print proceedings (no DOI)'}.get(r['kind'], 'no DOI')
        out.append(f'*{KIND_LABEL[r["kind"]]}* · `{key}` · {links}')
        out.append('')
    (OUT / 'references.md').write_text('\n'.join(out), encoding='utf-8')


# ------------------------------------------------------------------------------------------- brief (docx)
TOKENS = re.compile(r'(<i>|</i>|<b>|</b>|<sub>|</sub>)')


def add_runs(par, text, size=None, color=None):
    ital = bold = sub = False
    for tok in TOKENS.split(text):
        if tok == '<i>': ital = True
        elif tok == '</i>': ital = False
        elif tok == '<b>': bold = True
        elif tok == '</b>': bold = False
        elif tok == '<sub>': sub = True
        elif tok == '</sub>': sub = False
        elif tok:
            run = par.add_run(tok)
            run.italic = ital or None
            run.bold = bold or None
            if sub:
                run.font.subscript = True
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
    return par


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def add_page_number(section):
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tag:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), tag)
        else:
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = txt
        run._r.append(el)
    run.font.size = Pt(9)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin = sec.bottom_margin = Inches(0.65)
    st = doc.styles
    normal = st['Normal']
    normal.font.name = 'Calibri'
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before in (('Heading 1', 13, 12), ('Heading 2', 11, 8)):
        s = st[name]
        s.font.name = 'Calibri'
        rpr = s.element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
        for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
            rfonts.set(qn(attr), 'Calibri')
        for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.keep_with_next = True


def build_brief(n_comprehensive):
    import brief as B
    num = Numbering()
    doc = Document()
    style_document(doc)
    doc.core_properties.title = B.TITLE
    doc.core_properties.author = B.AUTHOR
    doc.core_properties.subject = B.SUBJECT
    doc.core_properties.keywords = 'sim-to-real; reality gap; UAV; PX4; DroneRanger; literature review'

    tp = doc.add_paragraph()
    add_runs(tp, f'<b>{B.TITLE}</b>', size=17)
    tp.paragraph_format.space_after = Pt(4)
    for i, line in enumerate(B.BYLINE):
        bp = doc.add_paragraph()
        add_runs(bp, line, size=10)
        bp.paragraph_format.space_after = Pt(1 if i < len(B.BYLINE) - 1 else 10)

    words = 0
    for kind, payload in B.CONTENT:
        if kind in ('h1', 'h2'):
            doc.add_paragraph(payload, style='Heading 1' if kind == 'h1' else 'Heading 2')
        elif kind == 'p':
            text = num.sub(payload.replace('{{N_COMPREHENSIVE}}', str(n_comprehensive)))
            if '{{' in text:
                raise SystemExit(f'unfilled placeholder in brief: {text[:80]}')
            words += len(re.sub(r'<[^>]+>', '', text).split())
            add_runs(doc.add_paragraph(), text)
        elif kind == 'entry':                         # annotated source: (key, lead-in, paragraph)
            key, lead, text = payload
            n = num.number(key)
            text = num.sub(text)
            words += len(re.sub(r'<[^>]+>', '', lead + ' ' + text).split())
            ep = doc.add_paragraph()
            add_runs(ep, f'<b>[{n}] {lead}</b> ' + text)
            ep.paragraph_format.left_indent = Inches(0.0)
        elif kind == 'bullets':
            for item in payload:
                item = num.sub(item)
                words += len(re.sub(r'<[^>]+>', '', item).split())
                bp = add_runs(doc.add_paragraph(style='List Bullet'), item)
                bp.paragraph_format.space_after = Pt(2)
        elif kind == 'table':
            caption, widths, header, rows = payload
            cp = add_runs(doc.add_paragraph(), f'<b>{caption}</b>', size=9.5)
            cp.paragraph_format.space_before = Pt(4); cp.paragraph_format.space_after = Pt(3)
            cp.paragraph_format.keep_with_next = True
            t = doc.add_table(rows=1, cols=len(header))
            t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            for r_i, cells in enumerate([header] + rows):
                row = t.rows[0] if r_i == 0 else t.add_row()
                tr_pr = row._tr.get_or_add_trPr()
                tr_pr.append(OxmlElement('w:cantSplit'))      # a row never breaks across pages
                if r_i == 0:
                    tr_pr.append(OxmlElement('w:tblHeader'))  # header row repeats if the table breaks
                for c_i, txt in enumerate(cells):
                    cell = row.cells[c_i]
                    cell.width = Inches(widths[c_i])
                    cpar = cell.paragraphs[0]
                    cpar.paragraph_format.space_after = Pt(0)
                    cpar.paragraph_format.line_spacing = 1.0
                    add_runs(cpar, f'<b>{txt}</b>' if r_i == 0 else num.sub(txt), size=8.5)
                    if r_i == 0:
                        shade(cell, 'D9D9D9')
                        cpar.paragraph_format.keep_with_next = True  # header stays with the first row
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        else:
            raise ValueError(kind)

    doc.add_paragraph('References', style='Heading 1')
    for n, key in enumerate(num.order, 1):
        rp = doc.add_paragraph()
        pf = rp.paragraph_format
        pf.left_indent = Inches(0.42); pf.first_line_indent = Inches(-0.42)
        pf.space_after = Pt(2); pf.tab_stops.add_tab_stop(Inches(0.42))
        pf.line_spacing = 1.0
        add_runs(rp, f'[{n}]\t' + ieee(BY_KEY[key]), size=9)
    add_page_number(doc.sections[0])
    out = OUT / 'DroneRanger_Sim2Real_Literature_Review_Brief.docx'
    doc.save(out)
    return num.order, words


def main():
    order = build_comprehensive()
    brief_order, brief_words = build_brief(len(order))
    build_reference_md(order, set(brief_order))
    write_bib(OUT / 'references.bib', order,
              '% Sim2Real literature review (DroneRanger), all cited references, in citation order.\n'
              '% Generated by _build/build.py from _build/refs.py on 2026-10-11. Import into Zotero: File > Import.')
    write_bib(OUT / 'references_brief.bib', brief_order,
              '% DroneRanger sim2real brief literature review: selected references (IEEE style in the .docx).\n'
              '% Generated by _build/build.py from _build/refs.py. Import into Zotero: File > Import.')
    missing = [k for k in brief_order if k not in order]
    unused = [r['key'] for r in REFS if r['key'] not in order and r['key'] not in brief_order]
    kinds = {}
    for k in brief_order:
        kinds[BY_KEY[k]['kind']] = kinds.get(BY_KEY[k]['kind'], 0) + 1
    src_words = len(CITE.sub('', review_source()).split())
    print(f'comprehensive: {len(order)} references, ~{src_words} words (incl. tables)')
    print(f'brief: {len(brief_order)} references by kind {kinds}; ~{brief_words} body words')
    print(f'brief refs not cited in comprehensive: {missing}')
    print(f'unused database entries: {unused}')
    req = kinds.get('J', 0) >= 6 and kinds.get('C', 0) >= 2 and kinds.get('B', 0) >= 1
    print('assignment counts (>=6 journal, >=2 conference, >=1 book):', 'OK' if req else 'NOT MET')
    for k in brief_order:
        if BY_KEY[k]['kind'] == 'J' and not BY_KEY[k].get('libkey'):
            print('WARNING: brief journal article without LibKey resolution:', k)


if __name__ == '__main__':
    main()
